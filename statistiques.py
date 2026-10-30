#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# =============================================================================
#  Statistiques — Raspberry Pi 5
#  Application de bureau (Tkinter) et cœur partagé avec telegram_bot_statistiques.py.
#  - Données : saisie / collage de texte, ou fichier (txt, csv, tsv, dat, xlsx)
#  - Calculs : délégués à R (moteur_statistique.R via cli_statistique.R)
#      statistiques descriptives, normalité (Shapiro-Wilk), tests t, ANOVA à un facteur
#      (+ Tukey), capabilité (Cp/Cpk/Pp/Ppk), Kappa de Cohen
#  - Au lancement : écran d'accueil avec icons/6Sigma.png
#
#  Dépendances : Python 3 + Tkinter (python3-tk) ; R + jsonlite (r-base, r-cran-jsonlite)
#                openpyxl (facultatif, lecture des fichiers .xlsx) ; Pillow (facultatif)
#  Lancement   : python3 statistiques.py
#  Le bot Telegram importe ce module : tkinter n'est chargé que pour l'interface.
#  Auteur : Jean-François BRUNET – JFBConseils – Novembre 2026
#  Rev.   :
# =============================================================================

import atexit
import csv
import json
import math
import os
import queue
import re
import shutil
import subprocess
import sys
import tempfile
import threading
from pathlib import Path

DOSSIER = Path(__file__).resolve().parent
ICONE = DOSSIER / "icons" / "6Sigma.png"
SCRIPT_R = DOSSIER / "cli_statistique.R"
TITRE = "Statistiques"
MAX_LIGNES = 200_000

MANQUANTS = {"", "na", "n/a", "nan", "null", "none"}
RE_DEC_VIRGULE = re.compile(r"^[-+]?\d+,\d+$")
RE_DEC_POINT = re.compile(r"^[-+]?\d+\.\d+$")
RE_ENTIER = re.compile(r"^[-+]?\d+$")


class ErreurStat(Exception):
    """Erreur destinée à l'utilisateur (message lisible)."""


# =============================================================================
#  Lecture des données
# =============================================================================
def _convertir(texte, dec):
    """(valeur, ok) : ok=False si la chaîne n'est pas un nombre ; valeur=None si manquante."""
    t = str(texte).strip().replace(" ", "").replace(" ", "")
    if t.lower() in MANQUANTS:
        return None, True
    if "_" in t:
        return None, False
    if dec == ",":
        t = t.replace(",", ".")
    try:
        v = float(t)
    except ValueError:
        return None, False
    if math.isnan(v) or math.isinf(v):
        return None, False
    return v, True


def _decouper(ligne, sep):
    if sep is None:
        return [ligne.strip()]
    if sep == " ":
        return ligne.split()
    try:
        return [c.strip() for c in next(csv.reader([ligne], delimiter=sep, skipinitialspace=True))]
    except Exception:
        return [c.strip() for c in ligne.split(sep)]


def _separateur(lignes):
    """Détecte le séparateur : '\\t', ';', ',', ' ' ou None (une seule colonne)."""
    ech = lignes[:50]
    moitie = len(ech) / 2
    for sep in ("\t", ";"):
        if sep in ech[0] and sum(sep in l for l in ech) >= moitie:
            return sep
    # Ni tabulation ni point-virgule : virgule décimale (données séparées par des espaces) ou virgule séparateur ?
    jetons = [t for l in ech for t in l.split()]
    if (jetons and any(RE_DEC_VIRGULE.match(t) for t in jetons)
            and all(RE_DEC_VIRGULE.match(t) or RE_ENTIER.match(t) for t in jetons)):
        return " " if any(len(l.split()) > 1 for l in ech) else None
    if "," in ech[0] and sum("," in l for l in ech) >= moitie:
        return ","
    if all(len(l.split()) > 1 for l in ech) and len({len(l.split()) for l in ech}) == 1:
        return " "
    if len(ech) == 1 and len(ech[0].split()) > 1:
        return " "
    return None


def _decimale(champs, sep):
    if sep == ",":
        return "."
    jetons = [c for ligne in champs for c in ligne]
    virg = sum(bool(RE_DEC_VIRGULE.match(t)) for t in jetons)
    pt = sum(bool(RE_DEC_POINT.match(t)) for t in jetons)
    return "," if virg > pt else "."


class Donnees:
    """Jeu de données : colonnes nommées, numériques (float / None) ou texte (str / None)."""

    def __init__(self, colonnes, types, avertissements=None, source="saisie"):
        self.colonnes = colonnes          # dict nom -> liste
        self.types = types                # dict nom -> "num" | "txt"
        self.avertissements = list(avertissements or [])
        self.source = source

    @property
    def noms(self):
        return list(self.colonnes)

    @property
    def noms_num(self):
        return [n for n in self.colonnes if self.types[n] == "num"]

    @property
    def noms_txt(self):
        return [n for n in self.colonnes if self.types[n] == "txt"]

    @property
    def nb_lignes(self):
        return max((len(v) for v in self.colonnes.values()), default=0)

    def resume(self):
        num = len(self.noms_num)
        return f"{self.nb_lignes} ligne(s), {len(self.noms)} colonne(s) dont {num} numérique(s)."

    def lignes(self, limite=None):
        n = self.nb_lignes if limite is None else min(limite, self.nb_lignes)
        cols = [self.colonnes[k] for k in self.noms]
        return [[(c[i] if i < len(c) else None) for c in cols] for i in range(n)]

    def ecrire_csv(self, chemin):
        with open(chemin, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            w.writerow(self.noms)
            for ligne in self.lignes():
                w.writerow(["" if v is None else (repr(v) if isinstance(v, float) else v) for v in ligne])

    def vers_csv_texte(self):
        sortie = []
        sortie.append(";".join(self.noms))
        for ligne in self.lignes():
            sortie.append(";".join("" if v is None else (fmt_nombre(v, 15) if isinstance(v, float) else str(v)) for v in ligne))
        return "\n".join(sortie) + "\n"


def lire_texte(texte, sep="auto", dec="auto", entete="auto"):
    """Interprète un texte (saisie ou contenu de fichier) : retourne un objet Donnees.

    sep    : 'auto', '\\t', ';', ',' ou ' '
    dec    : 'auto', '.' ou ','
    entete : 'auto', True ou False (la première ligne contient les noms de colonnes)
    """
    if texte is None:
        raise ErreurStat("Aucune donnée à lire.")
    lignes = [l for l in texte.replace("\r\n", "\n").replace("\r", "\n").replace("﻿", "").split("\n") if l.strip()]
    if not lignes:
        raise ErreurStat("Aucune donnée à lire.")
    if len(lignes) > MAX_LIGNES:
        raise ErreurStat(f"Trop de lignes (maximum {MAX_LIGNES}).")

    s = _separateur(lignes) if sep == "auto" else (None if sep in (None, "") else sep)
    champs = [_decouper(l, s) for l in lignes]
    d = _decimale(champs, s) if dec == "auto" else dec

    # Une seule ligne de nombres : c'est une série saisie à l'horizontale
    if len(champs) == 1 and entete in ("auto", False) and all(_convertir(c, d)[1] and c.strip() for c in champs[0]):
        champs = [[c] for c in champs[0]]
        entete = False

    premiere = champs[0]
    if entete == "auto":
        entete = any((not _convertir(c, d)[1]) for c in premiere)
    nc = max(len(r) for r in champs)
    if entete:
        noms = list(premiere) + [""] * (nc - len(premiere))
        corps = champs[1:]
    else:
        noms = [""] * nc
        corps = champs
    if not corps:
        raise ErreurStat("Aucune ligne de données : seule une ligne de noms de colonnes a été trouvée.")

    propres, vus = [], {}
    for i, n in enumerate(noms):
        n = n.strip().strip('"') or (f"Serie_{i + 1}" if not entete else f"Colonne_{i + 1}")
        base, k = n, 2
        while n in vus:
            n = f"{base}_{k}"
            k += 1
        vus[n] = True
        propres.append(n)

    colonnes, types, avert = {}, {}, []
    for j, nom in enumerate(propres):
        brut = [(r[j].strip() if j < len(r) else "") for r in corps]
        conv = [_convertir(v, d) for v in brut]
        non_manquantes = [v for v in brut if v.strip().lower() not in MANQUANTS]
        if not non_manquantes:
            avert.append(f"Colonne « {nom} » entièrement vide : ignorée.")
            continue
        if all(ok for _, ok in conv):
            colonnes[nom] = [v for v, _ in conv]
            types[nom] = "num"
        else:
            colonnes[nom] = [None if v.strip().lower() in MANQUANTS else v for v in brut]
            types[nom] = "txt"
            nb_num = sum(1 for (v, ok), b in zip(conv, brut) if ok and v is not None)
            if nb_num >= 0.5 * len(non_manquantes):
                mauvais = [b for (v, ok), b in zip(conv, brut) if not ok][:3]
                avert.append(f"Colonne « {nom} » : valeurs non numériques ({', '.join(mauvais)}…) — colonne traitée comme du texte.")
    if not colonnes:
        raise ErreurStat("Aucune colonne exploitable.")
    return Donnees(colonnes, types, avert)


def lire_fichier(chemin, sep="auto", dec="auto", entete="auto"):
    """Lit un fichier de données (txt, csv, tsv, dat, xlsx)."""
    p = Path(chemin)
    if not p.is_file():
        raise ErreurStat(f"Fichier introuvable : {p}")
    ext = p.suffix.lower()
    if ext == ".xlsx":
        return _lire_xlsx(p, entete)
    if ext in (".xls", ".ods", ".xlsm"):
        raise ErreurStat("Format non pris en charge : enregistrer le classeur en .xlsx ou en .csv.")
    brut = p.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texte = brut.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    d = lire_texte(texte, sep, dec, entete)
    d.source = p.name
    return d


def lire_texte_fichier(chemin):
    """Contenu texte d'un fichier (pour l'afficher dans la zone de saisie) ; None si illisible."""
    p = Path(chemin)
    if p.suffix.lower() in (".xlsx", ".xls", ".ods", ".xlsm") or p.stat().st_size > 1_000_000:
        return None
    brut = p.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return brut.decode(enc)
        except UnicodeDecodeError:
            continue
    return None


def _lire_xlsx(p, entete):
    try:
        import openpyxl
    except ImportError:
        raise ErreurStat("Lecture .xlsx indisponible : installer openpyxl (pip install openpyxl --break-system-packages) "
                         "ou enregistrer le classeur en .csv.")
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    lignes = []
    for row in ws.iter_rows(values_only=True):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        lignes.append("\t".join("" if c is None else str(c).replace("\t", " ").replace("\n", " ") for c in row))
    wb.close()
    d = lire_texte("\n".join(lignes), sep="\t", dec=".", entete=entete)
    d.source = p.name
    return d


# =============================================================================
#  Mise en forme
# =============================================================================
def fmt_nombre(v, chiffres=5):
    if v is None:
        return "—"
    if isinstance(v, bool):
        return "oui" if v else "non"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if math.isnan(v):
            return "—"
        if v == int(v) and abs(v) < 1e9:
            return str(int(v))
        return f"{v:.{chiffres}g}".replace(".", ",")
    return str(v)


LIBELLES = {"Serie": "Série", "Mediane": "Médiane", "Ecart_type": "Écart-type", "Ecart": "Écart", "P_Shapiro_Wilk": "p Shapiro-Wilk",
            "Sigma_global": "σ global", "Sigma_intra_fourni": "σ intra fourni", "Borne_inf": "Borne inf.", "Borne_sup": "Borne sup.",
            "p_value": "p-value", "p_ajustee": "p ajustée", "Somme_carres": "Somme des carrés", "Carre_moyen": "Carré moyen",
            "PPM_estimes_sous_LSL": "PPM estimés sous LSL", "PPM_estimes_sur_USL": "PPM estimés sur USL", "ddl": "ddl",
            "Reference": "Référence \\ Évaluation", "Series": "Séries"}


def libelle(c):
    return LIBELLES.get(c, str(c).replace("_", " "))


def fmt_cellule(nom, v):
    if nom == "Indicateur" and isinstance(v, str):
        return libelle(v)
    if isinstance(v, (int, float)) and not isinstance(v, bool) and str(nom).lower().startswith("p_"):
        return "< 0,0001" if v < 1e-4 else f"{v:.4f}".replace(".", ",")
    return fmt_nombre(v)


def table_texte(cols, rows, mode="large"):
    """Tableau en texte à chasse fixe (mode 'large') ou en fiches « Colonne : valeur » (mode 'fiche')."""
    cellules = [[fmt_cellule(c, v) for c, v in zip(cols, r)] for r in rows]
    if mode == "fiche" and list(cols) == ["Indicateur", "Valeur"]:
        return "\n".join(f"  {r[0]} : {r[1]}" for r in cellules)
    if mode == "fiche":
        blocs = []
        for r in cellules:
            blocs.append("\n".join(f"  {libelle(c)} : {v}" for c, v in zip(cols, r)))
        return "\n\n".join(blocs)
    ent = [libelle(c) for c in cols]
    larg = [max(len(ent[i]), *(len(r[i]) for r in cellules)) if cellules else len(ent[i]) for i in range(len(cols))]
    def ligne(vals):
        return "  ".join(v.ljust(larg[i]) if i == 0 else v.rjust(larg[i]) for i, v in enumerate(vals))
    sortie = [ligne(ent), "  ".join("-" * w for w in larg)] + [ligne(r) for r in cellules]
    return "\n".join(sortie)


def _liste(v):
    if v is None:
        return []
    if isinstance(v, (list, tuple)):
        return list(v)
    if isinstance(v, dict):
        return list(v.values())
    return [v]


class Resultat:
    """Résultat d'une analyse R : titre, lignes de texte, tableaux, figures PNG, remarques."""

    def __init__(self, brut, dossier):
        self.dossier = Path(dossier)
        self.commande = brut.get("commande", "")
        self.titre = brut.get("titre", "")
        self.lignes = [str(x) for x in _liste(brut.get("lignes"))]
        self.avertissements = [str(x) for x in _liste(brut.get("avertissements"))]
        self.figures = [self.dossier / str(f) for f in _liste(brut.get("figures")) if (self.dossier / str(f)).is_file()]
        self.tables = {}
        tabs = brut.get("tables")
        if isinstance(tabs, dict):
            for nom, t in tabs.items():
                if not isinstance(t, dict) or not t:
                    continue
                cols = list(t.keys())
                vals = [_liste(t[c]) for c in cols]
                self.tables[nom] = (cols, [list(r) for r in zip(*vals)])

    def texte(self, mode="large", titres=True):
        parts = []
        if titres and self.titre:
            parts.append(self.titre + "\n" + "=" * len(self.titre))
        if self.lignes:
            parts.append("\n".join(self.lignes))
        for nom, (cols, rows) in self.tables.items():
            parts.append(libelle(nom) + "\n" + table_texte(cols, rows, mode))
        if self.avertissements:
            parts.append("\n".join("Remarque : " + a for a in self.avertissements))
        return "\n\n".join(parts)

    def liberer(self):
        shutil.rmtree(self.dossier, ignore_errors=True)


# =============================================================================
#  Appel de R
# =============================================================================
_dossiers_temp = []


@atexit.register
def _nettoyer_temp():
    for d in _dossiers_temp:
        shutil.rmtree(d, ignore_errors=True)


def verifier_r():
    """(ok, message) : R et jsonlite sont-ils utilisables ?"""
    rs = shutil.which("Rscript")
    if not rs:
        return False, "R est introuvable. Installer : sudo apt install -y r-base r-cran-jsonlite"
    if not SCRIPT_R.is_file() or not (DOSSIER / "moteur_statistique.R").is_file():
        return False, "cli_statistique.R ou moteur_statistique.R manquant dans le dossier du programme."
    try:
        r = subprocess.run([rs, "-e", "library(jsonlite)"], capture_output=True, text=True, timeout=60)
    except Exception as e:
        return False, f"R inutilisable : {e}"
    if r.returncode != 0:
        return False, "Le package R « jsonlite » est absent. Installer : sudo apt install -y r-cran-jsonlite"
    return True, "R prêt."


def executer_r(donnees, commande, params, timeout=180):
    rs = shutil.which("Rscript")
    if not rs:
        raise ErreurStat("R est introuvable. Installer : sudo apt install -y r-base r-cran-jsonlite")
    dossier = Path(tempfile.mkdtemp(prefix="stat_"))
    _dossiers_temp.append(dossier)
    try:
        donnees.ecrire_csv(dossier / "donnees.csv")
        (dossier / "params.json").write_text(json.dumps({"commande": commande, **params}, ensure_ascii=False), encoding="utf-8")
        env = dict(os.environ, LC_ALL="C.UTF-8", LANG="C.UTF-8")
        try:
            p = subprocess.run([rs, str(SCRIPT_R), str(dossier / "donnees.csv"), str(dossier)],
                               capture_output=True, text=True, timeout=timeout, env=env)
        except subprocess.TimeoutExpired:
            raise ErreurStat("Le calcul R a dépassé le délai autorisé.")
        sortie = dossier / "resultat.json"
        if not sortie.is_file():
            detail = (p.stderr or p.stdout or "").strip().splitlines()
            raise ErreurStat("R n'a produit aucun résultat. " + (detail[-1] if detail else ""))
        brut = json.loads(sortie.read_text(encoding="utf-8"))
        if not brut.get("ok"):
            raise ErreurStat(str(brut.get("erreur", "Erreur de calcul.")))
        return Resultat(brut, dossier)
    except Exception:
        shutil.rmtree(dossier, ignore_errors=True)
        raise


def lire_nombre(texte, nom="valeur"):
    """Nombre saisi (virgule ou point décimal) ; chaîne vide -> None."""
    if texte is None:
        return None
    t = str(texte).strip().replace(" ", "")
    if t == "" or t == "-":
        return None
    try:
        v = float(t.replace(",", "."))
    except ValueError:
        raise ErreurStat(f"Valeur numérique invalide pour {nom} : « {texte} ».")
    if math.isnan(v) or math.isinf(v):
        raise ErreurStat(f"Valeur numérique invalide pour {nom} : « {texte} ».")
    return v


class Etude:
    """Jeu de données courant + analyses. Utilisé par l'interface et par le bot (une instance par utilisateur)."""

    def __init__(self):
        self.donnees = None
        self.echelle = 1.0                     # taille des graphiques R (1 = 900 px de large) ; réglée par l'interface

    def _r(self, commande, params):
        return executer_r(self._d(), commande, {**params, "echelle": self.echelle})

    def _d(self):
        if self.donnees is None:
            raise ErreurStat("Aucune donnée chargée.")
        return self.donnees

    def charger_texte(self, texte, **kw):
        self.donnees = lire_texte(texte, **kw)
        return self.donnees

    def charger_fichier(self, chemin, **kw):
        self.donnees = lire_fichier(chemin, **kw)
        return self.donnees

    def descriptif(self, colonnes=None):
        return self._r("descriptif", {"colonnes": list(colonnes) if colonnes else None})

    def normalite(self, colonne):
        return self._r("normalite", {"colonne": colonne})

    def test_t(self, x, y=None, mu=0.0, apparie=False, variances_egales=False, alternative="two.sided"):
        return self._r("t", {"x": x, "y": y or None, "mu": mu, "apparie": bool(apparie),
                             "variances_egales": bool(variances_egales), "alternative": alternative})

    def anova(self, mode, colonnes=None, valeur=None, groupe=None):
        return self._r("anova", {"mode": mode, "colonnes": list(colonnes) if colonnes else None,
                                 "valeur": valeur, "groupe": groupe})

    def capabilite(self, colonne, lsl=None, usl=None, cible=None, sigma_intra=None):
        if lsl is None and usl is None:
            raise ErreurStat("Saisir au moins une limite de spécification (LSL ou USL).")
        return self._r("capabilite", {"colonne": colonne, "lsl": lsl, "usl": usl,
                                      "cible": cible, "sigma_intra": sigma_intra})

    def kappa(self, a, b):
        return self._r("kappa", {"a": a, "b": b})


TEXTE_AIDE = """\
Statistiques — aide

Données :
• Une série de mesures par colonne ; la première ligne contient les noms de colonnes.
• Pour une analyse par groupes, utiliser une colonne de mesures et une colonne de groupes.
• Séparateurs reconnus : tabulation, point-virgule, virgule, espace. Décimale : point ou virgule.
• Les cases vides et les mentions NA sont traitées comme des valeurs manquantes.
• Fichiers : txt, csv, tsv, dat, xlsx. Les options de lecture permettent de forcer le séparateur,
  la décimale ou la présence d'une ligne de noms.
• Une saisie sans ligne de noms est possible si toutes les valeurs de la première ligne sont numériques.

Analyses : 
• Descriptif : effectif, moyenne, médiane, écart-type, variance, quartiles, extrêmes, Shapiro-Wilk.
• Normalité : test de Shapiro-Wilk, histogramme et droite de Henry.
• Test t : un échantillon, deux échantillons indépendants (Welch par défaut) ou appariés.
• ANOVA : un facteur, colonnes = groupes ou format long ; contrôle des résidus ;
  comparaisons de Tukey si le test global est significatif.
• Capabilité : Pp/Ppk à partir de l'écart-type global ; Cp/Cpk/Cm/Cmk uniquement si un écart-type
  intra-sous-groupe est fourni. Au moins une limite de spécification est requise.
• Kappa de Cohen : accord entre deux évaluations catégorielles (non pondéré).

Précautions : 
• Une p-value ne mesure ni l'importance pratique d'un effet ni la probabilité que l'hypothèse soit vraie.
• Un test de normalité non significatif ne prouve pas la normalité.
• Vérifier l'indépendance des observations, la stabilité du procédé et les unités avant toute décision.
• Les PPM reposent sur une loi normale ajustée : ils sont indicatifs.
"""


# =============================================================================
#  Interface Tkinter
# =============================================================================
def charger_image(chemin, largeur_max, hauteur_max=None, agrandir=False):
    """PhotoImage ajustée à largeur_max × hauteur_max (proportions conservées).

    Sans agrandir=True, l'image n'est jamais agrandie au-delà de sa taille d'origine.
    Pillow (avec ImageTk) donne un redimensionnement lissé ; sans lui, seule une réduction
    par pas entiers est possible (sudo apt install python3-pil.imagetk)."""
    import tkinter as tk
    try:
        from PIL import Image, ImageTk
        im = Image.open(chemin)
        ratio = largeur_max / im.width
        if hauteur_max:
            ratio = min(ratio, hauteur_max / im.height)
        if not agrandir:
            ratio = min(ratio, 1.0)
        if abs(ratio - 1.0) > 0.01:
            im = im.resize((max(1, int(im.width * ratio)), max(1, int(im.height * ratio))), Image.LANCZOS)
        return ImageTk.PhotoImage(im)
    except Exception:
        pass
    photo = tk.PhotoImage(file=str(chemin))
    ratio = largeur_max / photo.width()
    if hauteur_max:
        ratio = min(ratio, hauteur_max / photo.height())
    if ratio < 0.99:
        photo = photo.subsample(math.ceil(1 / ratio))
    return photo


def echelle_ecran(root):
    """Échelle de rendu des graphiques R (1 = 900 × 600 px), choisie pour que les graphiques restent nets
    en plein écran sur l'écran utilisé."""
    return max(1.0, min(2.0, root.winfo_screenwidth() * 0.96 / 900, root.winfo_screenheight() * 0.88 / 600))


class PanneauGraphiques:
    """Affiche des graphiques PNG ajustés à la place disponible ; ◀ ▶ pour passer de l'un à l'autre."""

    def __init__(self, parent, plein_ecran=None, fermer=None):
        import tkinter as tk
        from tkinter import ttk
        self.frame = ttk.Frame(parent)
        self.figures = []
        self.i = 0
        self.photo = None
        self._job = None
        self._taille = (0, 0)
        barre = ttk.Frame(self.frame)
        barre.pack(fill="x", pady=(0, 4))
        ttk.Button(barre, text="◀", width=3, command=lambda: self.aller(-1)).pack(side="left")
        self.lbl = ttk.Label(barre, text="", width=16, anchor="center")
        self.lbl.pack(side="left")
        ttk.Button(barre, text="▶", width=3, command=lambda: self.aller(1)).pack(side="left")
        ttk.Button(barre, text="Enregistrer ce graphique…", command=self.enregistrer).pack(side="left", padx=10)
        if plein_ecran:
            ttk.Button(barre, text="Plein écran", command=plein_ecran).pack(side="left")
        if fermer:
            ttk.Button(barre, text="Fermer", command=fermer).pack(side="right")
        self.zone = tk.Canvas(self.frame, background="white", highlightthickness=1, highlightbackground="#999999")
        self.zone.pack(fill="both", expand=True)
        self.zone.bind("<Configure>", self._reajuster)

    def definir(self, figures):
        self.figures = list(figures)
        self.i = 0
        self.afficher()

    def aller(self, pas):
        if len(self.figures) > 1:
            self.i = (self.i + pas) % len(self.figures)
            self.afficher()

    def afficher(self):
        self.zone.delete("all")
        self.photo = None
        w, h = self.zone.winfo_width(), self.zone.winfo_height()
        self._taille = (w, h)
        if not self.figures:
            self.lbl.configure(text="")
            return
        self.lbl.configure(text=f"Graphique {self.i + 1} / {len(self.figures)}")
        try:
            self.photo = charger_image(self.figures[self.i], max(200, w - 8), max(150, h - 8), agrandir=True)
            self.zone.create_image(w // 2, h // 2, image=self.photo)
        except Exception as e:
            self.zone.create_text(10, 10, anchor="nw", text=f"Affichage impossible : {e}")

    def _reajuster(self, _evt=None):
        if self._job:
            self.frame.after_cancel(self._job)
        self._job = self.frame.after(200, self._maj)

    def _maj(self):
        self._job = None
        w, h = self.zone.winfo_width(), self.zone.winfo_height()
        if self.figures and (abs(w - self._taille[0]) > 20 or abs(h - self._taille[1]) > 20):
            self.afficher()

    def enregistrer(self):
        from tkinter import filedialog
        if not self.figures:
            return
        src = self.figures[self.i]
        chemin = filedialog.asksaveasfilename(title="Enregistrer le graphique", defaultextension=".png", initialfile=src.name,
                                              filetypes=[("Image PNG", "*.png")], initialdir=str(Path.home()))
        if chemin:
            shutil.copy(src, chemin)


class Visionneuse:
    """Fenêtre plein écran pour voir les graphiques en grand."""

    def __init__(self, parent, figures):
        import tkinter as tk
        self.win = tk.Toplevel(parent)
        self.win.title(f"{TITRE} — graphiques")
        sw, sh = self.win.winfo_screenwidth(), self.win.winfo_screenheight()
        self.win.geometry(f"{int(sw * 0.96)}x{int(sh * 0.90)}+{int(sw * 0.02)}+0")
        self.panneau = PanneauGraphiques(self.win, fermer=self.win.destroy)
        self.panneau.frame.pack(fill="both", expand=True, padx=8, pady=8)
        self.win.bind("<Left>", lambda e: self.panneau.aller(-1))
        self.win.bind("<Right>", lambda e: self.panneau.aller(1))
        self.win.bind("<Escape>", lambda e: self.win.destroy())
        self.win.update_idletasks()
        self.panneau.definir(figures)


class ZoneResultat:
    """Zone de résultats : texte à chasse fixe à gauche, graphiques à droite (séparateur déplaçable)."""

    def __init__(self, parent, app):
        import tkinter as tk
        from tkinter import ttk, font as tkfont
        self.app = app
        self.resultat = None
        self.frame = ttk.Frame(parent)
        barre = ttk.Frame(self.frame)
        barre.pack(fill="x", pady=(6, 2))
        ttk.Button(barre, text="Enregistrer le résultat…", command=self.enregistrer).pack(side="left")
        ttk.Button(barre, text="Graphiques en plein écran…", command=self.agrandir).pack(side="left", padx=8)
        self.pan = ttk.PanedWindow(self.frame, orient="horizontal")
        self.pan.pack(fill="both", expand=True)
        gauche = ttk.Frame(self.pan)
        self.texte = tk.Text(gauche, wrap="none", height=10, width=40, state="disabled", padx=8, pady=6,
                             font=tkfont.nametofont("TkFixedFont"), background="white", relief="solid", borderwidth=1)
        sy = ttk.Scrollbar(gauche, orient="vertical", command=self.texte.yview)
        sx = ttk.Scrollbar(gauche, orient="horizontal", command=self.texte.xview)
        self.texte.configure(yscrollcommand=sy.set, xscrollcommand=sx.set)
        sx.pack(side="bottom", fill="x")
        sy.pack(side="right", fill="y")
        self.texte.pack(side="left", fill="both", expand=True)
        self.pan.add(gauche, weight=2)
        self.graphes = PanneauGraphiques(self.pan, plein_ecran=self.agrandir)
        self._graphes_visibles = False
        gras = tkfont.nametofont("TkFixedFont").copy()
        gras.configure(weight="bold")
        self.texte.tag_configure("titre", font=gras, foreground="#1f4e8c")
        self.texte.tag_configure("remarque", foreground="#8a5a00", wrap="word")
        self.texte.tag_configure("para", wrap="word")
        self.texte.tag_configure("erreur", foreground="#b00020")
        self.texte.tag_configure("section", font=gras)

    def _montrer_graphes(self, oui):
        if oui and not self._graphes_visibles:
            self.pan.add(self.graphes.frame, weight=3)
            self._graphes_visibles = True
            self.frame.after(80, self._placer_separateur)
            self.frame.after(500, lambda: self._placer_separateur(1))
        elif not oui and self._graphes_visibles:
            self.pan.forget(self.graphes.frame)
            self._graphes_visibles = False

    def _placer_separateur(self, essai=0):
        """Texte : 42 % de la largeur, graphiques : le reste. Réessaie tant que la zone n'est pas affichée ;
        ne touche pas à un séparateur déplacé à la main (sauf valeur aberrante)."""
        try:
            larg = self.pan.winfo_width()
            if larg < 200:
                raise ValueError("zone non affichée")
            pos = self.pan.sashpos(0)
            if essai == 0 or pos < 0.15 * larg or pos > 0.85 * larg:
                self.pan.sashpos(0, int(larg * 0.42))
        except Exception:
            if essai < 8:
                self.frame.after(150, lambda: self._placer_separateur(essai + 1))

    def agrandir(self):
        from tkinter import messagebox
        if self.resultat is None or not self.resultat.figures:
            messagebox.showinfo(TITRE, "Aucun graphique à agrandir.")
            return
        Visionneuse(self.app.root, self.resultat.figures)

    def _ecrire(self, contenu, tag=None):
        self.texte.insert("end", contenu, tag or ())

    def vider(self):
        self.texte.configure(state="normal")
        self.texte.delete("1.0", "end")
        self.texte.configure(state="disabled")

    def afficher_erreur(self, message):
        if self.resultat is not None:
            self.resultat.liberer()
            self.resultat = None
        self._montrer_graphes(False)
        self.graphes.definir([])
        self.vider()
        self.texte.configure(state="normal")
        self._ecrire("Erreur : " + message + "\n", "erreur")
        self.texte.configure(state="disabled")

    def afficher(self, res):
        if self.resultat is not None and self.resultat is not res:
            self.resultat.liberer()
        self.resultat = res
        self.vider()
        t = self.texte
        t.configure(state="normal")
        self._ecrire(res.titre + "\n", "titre")
        self._ecrire("=" * len(res.titre) + "\n\n")
        if res.lignes:
            self._ecrire("\n".join(res.lignes) + "\n\n", "para")
        for nom, (cols, rows) in res.tables.items():
            self._ecrire(libelle(nom) + "\n", "section")
            self._ecrire(table_texte(cols, rows) + "\n\n")
        for a in res.avertissements:
            self._ecrire("Remarque : " + a + "\n", "remarque")
        t.configure(state="disabled")
        t.yview_moveto(0)
        self._montrer_graphes(bool(res.figures))
        self.graphes.definir(res.figures)

    def enregistrer(self):
        from tkinter import filedialog, messagebox
        if self.resultat is None:
            messagebox.showinfo(TITRE, "Aucun résultat à enregistrer.")
            return
        chemin = filedialog.asksaveasfilename(title="Enregistrer le résultat", defaultextension=".txt",
                                              filetypes=[("Texte", "*.txt")], initialdir=str(Path.home()))
        if not chemin:
            return
        p = Path(chemin)
        p.write_text(self.resultat.texte() + "\n", encoding="utf-8")
        for f in self.resultat.figures:
            shutil.copy(f, p.with_name(f"{p.stem}_{f.name}"))
        self.app.statut(f"Résultat enregistré : {p.name}" + (" (figures PNG à côté)" if self.resultat.figures else ""))


class Application:
    def __init__(self, root):
        import tkinter as tk
        from tkinter import ttk
        self.tk, self.ttk = tk, ttk
        self.root = root
        self.etude = Etude()
        self.file = queue.Queue()
        self.occupe = False
        root.title(f"{TITRE} — JFBConseils")
        sw, sh = root.winfo_screenwidth(), root.winfo_screenheight()
        root.geometry(f"{int(sw * 0.96)}x{int(sh * 0.92)}+{int(sw * 0.02)}+0")   # grande fenêtre, adaptée à l'écran
        try:
            root.attributes("-zoomed", True)                                      # agrandie au maximum (X11)
        except Exception:
            pass
        root.minsize(860, 600)
        self.statut_var = tk.StringVar(value="Prêt.")
        self._construire()
        self._maj_choix()
        root.protocol("WM_DELETE_WINDOW", self.quitter)
        root.after(120, self._surveiller)
        threading.Thread(target=self._verifier_r, daemon=True).start()

    # -- utilitaires ------------------------------------------------------------------
    def statut(self, texte):
        self.statut_var.set(texte)

    def _verifier_r(self):
        ok, msg = verifier_r()
        if not ok:
            self.file.put(("statut", "⚠ " + msg))

    def _surveiller(self):
        try:
            while True:
                item = self.file.get_nowait()
                if item[0] == "statut":
                    self.statut(item[1])
                    continue
                zone, res, err = item
                self.occupe = False
                if err:
                    zone.afficher_erreur(err)
                    self.statut("Erreur.")
                else:
                    zone.afficher(res)
                    self.statut("Calcul terminé.")
        except queue.Empty:
            pass
        self.root.after(120, self._surveiller)

    def lancer(self, zone, fonction):
        if self.etude.donnees is None:
            zone.afficher_erreur("Aucune donnée chargée : utiliser l'onglet « Données ».")
            return
        if self.occupe:
            self.statut("Un calcul est déjà en cours…")
            return
        self.occupe = True
        self.etude.echelle = echelle_ecran(self.root)                    # graphiques R nets à la taille de l'écran
        self.statut("Calcul en cours…")

        def tache():
            try:
                self.file.put((zone, fonction(), None))
            except ErreurStat as e:
                self.file.put((zone, None, str(e)))
            except Exception as e:
                self.file.put((zone, None, f"{type(e).__name__} : {e}"))
        threading.Thread(target=tache, daemon=True).start()

    def quitter(self):
        for z in self.zones:
            if z.resultat is not None:
                z.resultat.liberer()
        self.root.destroy()

    # -- construction -----------------------------------------------------------------
    def _construire(self):
        ttk = self.ttk
        self.zones = []
        nb = ttk.Notebook(self.root)
        nb.pack(fill="both", expand=True, padx=8, pady=(8, 0))
        ttk.Label(self.root, textvariable=self.statut_var, anchor="w", relief="sunken", padding=(6, 2)).pack(fill="x", padx=8, pady=6)
        self.nb = nb
        self._onglet_donnees(nb)
        self._onglet_descriptif(nb)
        self._onglet_normalite(nb)
        self._onglet_t(nb)
        self._onglet_anova(nb)
        self._onglet_capabilite(nb)
        self._onglet_kappa(nb)
        self._onglet_aide(nb)

    def _nouvelle_zone(self, parent):
        z = ZoneResultat(parent, self)
        z.frame.pack(fill="both", expand=True, padx=8, pady=(0, 8))
        self.zones.append(z)
        return z

    def _onglet(self, nb, titre):
        f = self.ttk.Frame(nb)
        nb.add(f, text=titre)
        return f

    # --- Données
    def _onglet_donnees(self, nb):
        tk, ttk = self.tk, self.ttk
        f = self._onglet(nb, "Données")
        barre = ttk.Frame(f)
        barre.pack(fill="x", padx=8, pady=8)
        ttk.Button(barre, text="Ouvrir un fichier…", command=self.ouvrir_fichier).pack(side="left", padx=(0, 6))
        ttk.Button(barre, text="Charger la saisie", command=self.charger_saisie).pack(side="left", padx=6)
        ttk.Button(barre, text="Effacer la saisie", command=lambda: self.saisie.delete("1.0", "end")).pack(side="left", padx=6)
        ttk.Button(barre, text="Exporter en CSV…", command=self.exporter_csv).pack(side="left", padx=6)

        opt = ttk.LabelFrame(f, text="Lecture")
        opt.pack(fill="x", padx=8)
        self.v_sep = tk.StringVar(value="Automatique")
        self.v_dec = tk.StringVar(value="Automatique")
        self.v_ent = tk.StringVar(value="Automatique")
        for i, (lib, var, vals) in enumerate([
                ("Séparateur", self.v_sep, ["Automatique", "Tabulation", "Point-virgule", "Virgule", "Espace"]),
                ("Décimale", self.v_dec, ["Automatique", "Point", "Virgule"]),
                ("Ligne de noms de colonnes", self.v_ent, ["Automatique", "Oui", "Non"])]):
            ttk.Label(opt, text=lib).grid(row=0, column=2 * i, padx=(8, 4), pady=6, sticky="e")
            ttk.Combobox(opt, textvariable=var, values=vals, state="readonly", width=15).grid(row=0, column=2 * i + 1, padx=(0, 8), pady=6)

        pan = ttk.PanedWindow(f, orient="vertical")
        pan.pack(fill="both", expand=True, padx=8, pady=8)
        cadre_s = ttk.LabelFrame(pan, text="Saisie ou collage (une série par colonne, noms de colonnes en première ligne)")
        from tkinter import font as tkfont
        self.saisie = tk.Text(cadre_s, height=8, undo=True, wrap="none", font=tkfont.nametofont("TkFixedFont"))
        sy = ttk.Scrollbar(cadre_s, orient="vertical", command=self.saisie.yview)
        self.saisie.configure(yscrollcommand=sy.set)
        sy.pack(side="right", fill="y")
        self.saisie.pack(fill="both", expand=True, padx=4, pady=4)
        pan.add(cadre_s, weight=1)

        cadre_a = ttk.LabelFrame(pan, text="Données chargées")
        self.resume_var = tk.StringVar(value="Aucune donnée chargée.")
        ttk.Label(cadre_a, textvariable=self.resume_var).pack(anchor="w", padx=6, pady=4)
        corps = ttk.Frame(cadre_a)
        corps.pack(fill="both", expand=True, padx=4, pady=(0, 4))
        self.arbre = ttk.Treeview(corps, show="headings", height=8)
        ay = ttk.Scrollbar(corps, orient="vertical", command=self.arbre.yview)
        ax = ttk.Scrollbar(corps, orient="horizontal", command=self.arbre.xview)
        self.arbre.configure(yscrollcommand=ay.set, xscrollcommand=ax.set)
        ax.pack(side="bottom", fill="x")
        ay.pack(side="right", fill="y")
        self.arbre.pack(side="left", fill="both", expand=True)
        pan.add(cadre_a, weight=2)

    def _options_lecture(self):
        sep = {"Automatique": "auto", "Tabulation": "\t", "Point-virgule": ";", "Virgule": ",", "Espace": " "}[self.v_sep.get()]
        dec = {"Automatique": "auto", "Point": ".", "Virgule": ","}[self.v_dec.get()]
        ent = {"Automatique": "auto", "Oui": True, "Non": False}[self.v_ent.get()]
        return dict(sep=sep, dec=dec, entete=ent)

    def charger_saisie(self):
        from tkinter import messagebox
        try:
            self.etude.charger_texte(self.saisie.get("1.0", "end"), **self._options_lecture())
        except ErreurStat as e:
            messagebox.showerror(TITRE, str(e))
            return
        self._donnees_chargees("saisie")

    def ouvrir_fichier(self):
        from tkinter import filedialog, messagebox
        chemin = filedialog.askopenfilename(title="Ouvrir un fichier de données", initialdir=str(Path.home()),
                                            filetypes=[("Fichiers de données", "*.csv *.txt *.tsv *.dat *.xlsx"), ("Tous les fichiers", "*.*")])
        if not chemin:
            return
        try:
            self.etude.charger_fichier(chemin, **self._options_lecture())
            contenu = lire_texte_fichier(chemin)
        except ErreurStat as e:
            messagebox.showerror(TITRE, str(e))
            return
        except Exception as e:
            messagebox.showerror(TITRE, f"Lecture impossible : {e}")
            return
        if contenu is not None:
            self.saisie.delete("1.0", "end")
            self.saisie.insert("1.0", contenu)
        self._donnees_chargees(Path(chemin).name)

    def _donnees_chargees(self, origine):
        d = self.etude.donnees
        self.resume_var.set(f"{origine} — {d.resume()}")
        arbre = self.arbre
        arbre.delete(*arbre.get_children())
        arbre["columns"] = d.noms
        for n in d.noms:
            arbre.heading(n, text=n + ("" if d.types[n] == "num" else " (texte)"))
            arbre.column(n, width=110, anchor="e" if d.types[n] == "num" else "w", stretch=False)
        for ligne in d.lignes(500):
            arbre.insert("", "end", values=[fmt_nombre(v, 10) if v is not None else "" for v in ligne])
        self._maj_choix()
        msg = d.resume()
        if d.avertissements:
            msg += " Attention : " + " ".join(d.avertissements)
        self.statut(msg)

    def exporter_csv(self):
        from tkinter import filedialog, messagebox
        if self.etude.donnees is None:
            messagebox.showinfo(TITRE, "Aucune donnée chargée.")
            return
        chemin = filedialog.asksaveasfilename(title="Exporter les données", defaultextension=".csv",
                                              filetypes=[("CSV", "*.csv")], initialdir=str(Path.home()))
        if chemin:
            Path(chemin).write_text(self.etude.donnees.vers_csv_texte(), encoding="utf-8-sig")
            self.statut(f"Données exportées : {Path(chemin).name} (séparateur point-virgule, décimale virgule)")

    # --- Onglets d'analyse
    def _liste_multi(self, parent, hauteur=4):
        tk, ttk = self.tk, self.ttk
        cadre = ttk.Frame(parent)
        lb = tk.Listbox(cadre, selectmode="multiple", exportselection=False, height=hauteur, width=28)
        sb = ttk.Scrollbar(cadre, orient="vertical", command=lb.yview)
        lb.configure(yscrollcommand=sb.set)
        lb.pack(side="left", fill="both")
        sb.pack(side="left", fill="y")
        return cadre, lb

    def _combo(self, parent, var, largeur=26):
        return self.ttk.Combobox(parent, textvariable=var, state="readonly", width=largeur, values=[])

    def _onglet_descriptif(self, nb):
        ttk = self.ttk
        f = self._onglet(nb, "Descriptif")
        c = ttk.Frame(f)
        c.pack(fill="x", padx=8, pady=8)
        ttk.Label(c, text="Séries à analyser").grid(row=0, column=0, sticky="nw", padx=(0, 8))
        cadre, self.lb_desc = self._liste_multi(c)
        cadre.grid(row=0, column=1, sticky="w")
        b = ttk.Frame(c)
        b.grid(row=0, column=2, padx=12, sticky="n")
        ttk.Button(b, text="Tout sélectionner", command=lambda: self.lb_desc.select_set(0, "end")).pack(fill="x", pady=2)
        ttk.Button(b, text="Lancer", command=self.run_descriptif).pack(fill="x", pady=2)
        self.z_desc = self._nouvelle_zone(f)

    def run_descriptif(self):
        cols = [self.lb_desc.get(i) for i in self.lb_desc.curselection()] or None
        self.lancer(self.z_desc, lambda: self.etude.descriptif(cols))

    def _onglet_normalite(self, nb):
        tk, ttk = self.tk, self.ttk
        f = self._onglet(nb, "Normalité")
        c = ttk.Frame(f)
        c.pack(fill="x", padx=8, pady=8)
        self.v_norm = tk.StringVar()
        ttk.Label(c, text="Série").pack(side="left", padx=(0, 8))
        self.cb_norm = self._combo(c, self.v_norm)
        self.cb_norm.pack(side="left")
        ttk.Button(c, text="Lancer", command=self.run_normalite).pack(side="left", padx=12)
        self.z_norm = self._nouvelle_zone(f)

    def run_normalite(self):
        col = self.v_norm.get()                      # lecture des variables Tk dans le thread principal
        self.lancer(self.z_norm, lambda: self.etude.normalite(col))

    def _onglet_t(self, nb):
        tk, ttk = self.tk, self.ttk
        f = self._onglet(nb, "Test t")
        c = ttk.Frame(f)
        c.pack(fill="x", padx=8, pady=8)
        self.v_tx, self.v_ty = tk.StringVar(), tk.StringVar()
        self.v_mu = tk.StringVar(value="0")
        self.v_talt = tk.StringVar(value="Différent de")
        self.v_tapp, self.v_tvar = tk.BooleanVar(value=False), tk.BooleanVar(value=False)
        ttk.Label(c, text="Série X").grid(row=0, column=0, sticky="e", padx=4, pady=3)
        self.cb_tx = self._combo(c, self.v_tx)
        self.cb_tx.grid(row=0, column=1, sticky="w", pady=3)
        ttk.Label(c, text="Série Y").grid(row=1, column=0, sticky="e", padx=4, pady=3)
        self.cb_ty = self._combo(c, self.v_ty)
        self.cb_ty.grid(row=1, column=1, sticky="w", pady=3)
        ttk.Label(c, text="(« — un échantillon — » : comparaison de X à une valeur de référence)").grid(row=1, column=2, sticky="w", padx=8)
        ttk.Label(c, text="Valeur de référence").grid(row=2, column=0, sticky="e", padx=4, pady=3)
        ttk.Entry(c, textvariable=self.v_mu, width=14).grid(row=2, column=1, sticky="w", pady=3)
        ttk.Label(c, text="Hypothèse alternative").grid(row=3, column=0, sticky="e", padx=4, pady=3)
        ttk.Combobox(c, textvariable=self.v_talt, state="readonly", width=24,
                     values=["Différent de", "Supérieur à", "Inférieur à"]).grid(row=3, column=1, sticky="w", pady=3)
        ttk.Checkbutton(c, text="Mesures appariées", variable=self.v_tapp).grid(row=4, column=1, sticky="w", pady=2)
        ttk.Checkbutton(c, text="Variances supposées égales (échantillons indépendants)", variable=self.v_tvar).grid(row=5, column=1, sticky="w", pady=2)
        ttk.Button(c, text="Lancer", command=self.run_t).grid(row=6, column=1, sticky="w", pady=8)
        self.z_t = self._nouvelle_zone(f)

    def run_t(self):
        try:
            mu = lire_nombre(self.v_mu.get(), "la valeur de référence") or 0.0
        except ErreurStat as e:
            self.z_t.afficher_erreur(str(e))
            return
        y = self.v_ty.get()
        y = "" if y.startswith("—") else y
        alt = {"Différent de": "two.sided", "Supérieur à": "greater", "Inférieur à": "less"}[self.v_talt.get()]
        x = self.v_tx.get()
        apparie, var_egales = self.v_tapp.get(), self.v_tvar.get()
        self.lancer(self.z_t, lambda: self.etude.test_t(x, y or None, mu, apparie, var_egales, alt))

    def _onglet_anova(self, nb):
        tk, ttk = self.tk, self.ttk
        f = self._onglet(nb, "ANOVA")
        c = ttk.Frame(f)
        c.pack(fill="x", padx=8, pady=8)
        self.v_amode = tk.StringVar(value="large")
        ttk.Radiobutton(c, text="Chaque colonne numérique est un groupe", variable=self.v_amode, value="large").grid(row=0, column=0, columnspan=3, sticky="w")
        cadre, self.lb_anova = self._liste_multi(c)
        cadre.grid(row=1, column=0, columnspan=3, sticky="w", padx=20, pady=4)
        ttk.Radiobutton(c, text="Une colonne de mesures et une colonne de groupes", variable=self.v_amode, value="long").grid(row=2, column=0, columnspan=3, sticky="w", pady=(8, 0))
        self.v_aval, self.v_agrp = tk.StringVar(), tk.StringVar()
        ttk.Label(c, text="Mesures").grid(row=3, column=0, sticky="e", padx=(20, 4), pady=3)
        self.cb_aval = self._combo(c, self.v_aval)
        self.cb_aval.grid(row=3, column=1, sticky="w")
        ttk.Label(c, text="Groupes").grid(row=4, column=0, sticky="e", padx=(20, 4), pady=3)
        self.cb_agrp = self._combo(c, self.v_agrp)
        self.cb_agrp.grid(row=4, column=1, sticky="w")
        ttk.Button(c, text="Lancer", command=self.run_anova).grid(row=5, column=1, sticky="w", pady=8)
        self.z_anova = self._nouvelle_zone(f)

    def run_anova(self):
        if self.v_amode.get() == "large":
            cols = [self.lb_anova.get(i) for i in self.lb_anova.curselection()]
            self.lancer(self.z_anova, lambda: self.etude.anova("large", colonnes=cols))
        else:
            v, g = self.v_aval.get(), self.v_agrp.get()
            self.lancer(self.z_anova, lambda: self.etude.anova("long", valeur=v, groupe=g))

    def _onglet_capabilite(self, nb):
        tk, ttk = self.tk, self.ttk
        f = self._onglet(nb, "Capabilité")
        c = ttk.Frame(f)
        c.pack(fill="x", padx=8, pady=8)
        self.v_cs = tk.StringVar()
        self.v_lsl, self.v_cib, self.v_usl, self.v_sig = (tk.StringVar() for _ in range(4))
        ttk.Label(c, text="Série de mesures").grid(row=0, column=0, sticky="e", padx=4, pady=3)
        self.cb_cap = self._combo(c, self.v_cs)
        self.cb_cap.grid(row=0, column=1, sticky="w")
        for i, (lib, var) in enumerate([("LSL — limite inférieure", self.v_lsl), ("Cible (facultatif)", self.v_cib),
                                        ("USL — limite supérieure", self.v_usl),
                                        ("Écart-type intra-sous-groupe (facultatif)", self.v_sig)], start=1):
            ttk.Label(c, text=lib).grid(row=i, column=0, sticky="e", padx=4, pady=3)
            ttk.Entry(c, textvariable=var, width=16).grid(row=i, column=1, sticky="w")
        ttk.Label(c, text="Laisser vide une limite absente. Sans écart-type intra-sous-groupe, seuls Pp et Ppk sont calculés.").grid(row=5, column=0, columnspan=3, sticky="w", padx=4, pady=3)
        ttk.Button(c, text="Calculer", command=self.run_cap).grid(row=6, column=1, sticky="w", pady=8)
        self.z_cap = self._nouvelle_zone(f)

    def run_cap(self):
        try:
            lsl = lire_nombre(self.v_lsl.get(), "LSL")
            usl = lire_nombre(self.v_usl.get(), "USL")
            cib = lire_nombre(self.v_cib.get(), "la cible")
            sig = lire_nombre(self.v_sig.get(), "l'écart-type intra-sous-groupe")
        except ErreurStat as e:
            self.z_cap.afficher_erreur(str(e))
            return
        col = self.v_cs.get()
        self.lancer(self.z_cap, lambda: self.etude.capabilite(col, lsl, usl, cib, sig))

    def _onglet_kappa(self, nb):
        tk, ttk = self.tk, self.ttk
        f = self._onglet(nb, "Kappa")
        c = ttk.Frame(f)
        c.pack(fill="x", padx=8, pady=8)
        self.v_ka, self.v_kb = tk.StringVar(), tk.StringVar()
        ttk.Label(c, text="Évaluation de référence").grid(row=0, column=0, sticky="e", padx=4, pady=3)
        self.cb_ka = self._combo(c, self.v_ka)
        self.cb_ka.grid(row=0, column=1, sticky="w")
        ttk.Label(c, text="Deuxième évaluation").grid(row=1, column=0, sticky="e", padx=4, pady=3)
        self.cb_kb = self._combo(c, self.v_kb)
        self.cb_kb.grid(row=1, column=1, sticky="w")
        ttk.Button(c, text="Calculer", command=self.run_kappa).grid(row=2, column=1, sticky="w", pady=8)
        self.z_kappa = self._nouvelle_zone(f)

    def run_kappa(self):
        a, b = self.v_ka.get(), self.v_kb.get()
        self.lancer(self.z_kappa, lambda: self.etude.kappa(a, b))

    def _onglet_aide(self, nb):
        tk = self.tk
        from tkinter import font as tkfont
        f = self._onglet(nb, "Aide")
        t = tk.Text(f, wrap="word", padx=12, pady=10, background="white", relief="flat", font=tkfont.nametofont("TkDefaultFont"))
        t.insert("1.0", TEXTE_AIDE)
        t.configure(state="disabled")
        t.pack(fill="both", expand=True, padx=8, pady=8)

    # -- mise à jour des listes de colonnes -----------------------------------------------
    def _maj_choix(self):
        d = self.etude.donnees
        num = d.noms_num if d else []
        tout = d.noms if d else []
        txt = d.noms_txt if d else []
        self.lb_desc.delete(0, "end")
        self.lb_anova.delete(0, "end")
        for n in num:
            self.lb_desc.insert("end", n)
            self.lb_anova.insert("end", n)
        self.lb_desc.select_set(0, "end")
        self.lb_anova.select_set(0, "end")
        for cb, vals, var in [(self.cb_norm, num, self.v_norm), (self.cb_tx, num, self.v_tx), (self.cb_cap, num, self.v_cs),
                              (self.cb_aval, num, self.v_aval), (self.cb_agrp, tout, self.v_agrp),
                              (self.cb_ka, tout, self.v_ka), (self.cb_kb, tout, self.v_kb)]:
            cb["values"] = vals
            var.set(vals[0] if vals else "")
        self.cb_ty["values"] = ["— un échantillon —"] + num
        self.v_ty.set("— un échantillon —")
        if len(num) >= 2:
            self.v_ty.set(num[1])
        cand = txt if len(txt) >= 2 else tout          # Kappa : colonnes texte en priorité
        if len(cand) >= 2:
            self.v_ka.set(cand[0])
            self.v_kb.set(cand[1])
        if txt and num:                                # colonne de groupes textuelle : format long par défaut
            self.v_amode.set("long")
            self.v_agrp.set(txt[0])
        else:
            self.v_amode.set("large")
            if len(num) >= 2:
                self.v_agrp.set(num[1])


class EcranAccueil:
    """Fenêtre d'accueil sans bordure avec l'icône du projet."""

    def __init__(self, root, duree_ms, suite):
        import tkinter as tk
        self.win = tk.Toplevel(root)
        self.win.overrideredirect(True)
        try:
            self.win.attributes("-topmost", True)
        except Exception:
            pass
        cadre = tk.Frame(self.win, bg="white", highlightbackground="#1f4e8c", highlightthickness=2)
        cadre.pack(fill="both", expand=True)
        self.photo = None
        if ICONE.is_file():
            try:
                self.photo = charger_image(ICONE, 340)
                tk.Label(cadre, image=self.photo, bg="white").pack(padx=30, pady=(24, 8))
            except Exception:
                self.photo = None
        tk.Label(cadre, text=TITRE, bg="white", fg="#1f4e8c", font=("TkDefaultFont", 20, "bold")).pack(padx=30, pady=(8 if self.photo else 30, 2))
        tk.Label(cadre, text="JFBConseils — Statistiques", bg="white", fg="#555555").pack()
        tk.Label(cadre, text="Chargement…", bg="white", fg="#888888").pack(pady=(10, 22))
        self.win.update_idletasks()
        w, h = self.win.winfo_reqwidth(), self.win.winfo_reqheight()
        x = (self.win.winfo_screenwidth() - w) // 2
        y = (self.win.winfo_screenheight() - h) // 2
        self.win.geometry(f"{w}x{h}+{x}+{y}")
        self.win.update()
        root.after(duree_ms, lambda: (self.win.destroy(), suite()))


def main():
    try:
        import tkinter as tk
    except ImportError:
        sys.exit("Tkinter est absent : sudo apt install -y python3-tk")
    import tkinter.font  # noqa: F401  (rend tk.font disponible)
    root = tk.Tk()
    root.withdraw()
    if ICONE.is_file():
        try:
            root.iconphoto(True, charger_image(ICONE, 128))
        except Exception:
            pass

    def ouvrir():
        Application(root)
        root.deiconify()
    EcranAccueil(root, 1800, ouvrir)
    root.mainloop()


if __name__ == "__main__":
    main()
